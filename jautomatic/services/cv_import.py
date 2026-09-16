"""Extract a Profile from a CV file (docx, txt or md).

Fully offline and deterministic: the wrapper format is stripped to plain text,
then structural heuristics (contact regexes, section headings, date ranges and
list bullets) fill as much of the Profile as possible.  CVs vary wildly, so the
result is a draft for the Profile tab — reviewed and edited by the user before
anything is generated.
"""
from __future__ import annotations

import re
from pathlib import Path

from ..models import (
    GENERIC_TERMS,
    STOPWORDS,
    EducationEntry,
    ExperienceEntry,
    Profile,
    pretty_term,
)

SUPPORTED_SUFFIXES = (".docx", ".txt", ".md")

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
URL_RE = re.compile(r"(?:https?://|www\.)[^\s)\]\",;]+", re.IGNORECASE)
LINK_RE = re.compile(r"(?:linkedin|github)\.com/[^\s)\]\",;]+", re.IGNORECASE)
PHONE_RE = re.compile(r"(?<!\d)\+?\d[\d\s().-]{6,}\d")
_MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?"
RANGE_RE = re.compile(
    r"\b(" + _MONTH + r"\.?\s*)?" + r"((?:19|20)\d{2})"
    r"\s*(?:[-–—~]|to|until)\s*"
    r"(" + _MONTH + r"\.?\s*)?(present|current|now|(?:19|20)\d{2})\b",
    re.IGNORECASE,
)
_NAME_RE = re.compile(
    r"^([A-ZÀ-Ž][A-Za-zÀ-ž.'-]+|[A-ZÀ-Ž]\.)\s+"
    r"([A-ZÀ-Ž][A-Za-zÀ-ž.'-]+|[A-ZÀ-Ž]\.)"
    r"(\s+[A-ZÀ-Ž][A-Za-zÀ-ž.'-]+){0,2}$"
)
_CITY_RE = re.compile(r"^[A-ZÀ-Ž][A-Za-zÀ-ž .'\-]{2,},\s*[A-ZÀ-Ž][A-Za-zÀ-ž .'\-]{2,}$")
_LOCATION_LABEL_RE = re.compile(r"^(?:location|city|address|based in)\s*[:\-]\s*(.+)$",
                                re.IGNORECASE)
_BULLET_RE = re.compile(r"^\s*[-–—•·*+]\s+|^\s*\d+[.)]\s+|^\s*>[>\s]*\s+")
_DEGREE_RE = re.compile(
    r"\b(?:bachelor|b\.?s\.?c|master|m\.?s\.?c|mba|ph\.?d\.?|doctor|b\.?a\.?|m\.?a\.?|"
    r"diploma|certificate|degree|course|abitur|baccalaur[eé]at)\b",
    re.IGNORECASE,
)
_SCHOOL_RE = re.compile(
    r"\b(?:university|college|institute|academy|school|polytechnic|gymnasium|facultad)\b",
    re.IGNORECASE,
)

_SECTION_PATTERNS = {
    "summary": re.compile(r"^(?:summary|professional summary|profile|about me|about|objective|"
                          r"career objective)\b", re.IGNORECASE),
    "skills": re.compile(r"^(?:skills?\b|key skills?\b|core skills?\b|skills summary|"
                         r"skill set|technical skills|professional skills|digital skills|"
                         r"computer skills|it skills|technologies|tech stack|"
                         r"technical competencies|core competencies|competencies|"
                         r"areas of expertise|expertise|technical strengths)\b", re.IGNORECASE),
    "experience": re.compile(r"^(?:experience|relevant experience|work experience|employment|"
                             r"employment history|professional experience|work history|"
                             r"career history|professional background|career)\b", re.IGNORECASE),
    "education": re.compile(r"^(?:education|academic background|academic|qualifications)\b",
                            re.IGNORECASE),
    "languages": re.compile(r"^languages?\b", re.IGNORECASE),
}

_STRIP_MARKERS_RE = re.compile(r"^[\s#>=\-*•·*]+")

_SINGLE_DATE_RE = re.compile(
    r"^\s*(?:" + _MONTH + r"\.?\s+)?(?:19|20)\d{2}(?:\s*[-/]\s*\d{1,2})?\s*$",
    re.IGNORECASE,
)
_SKILL_LABELS = {
    "languages", "language", "spoken languages", "programming languages", "programming",
    "technologies", "technology", "tech stack", "technical skills", "core competencies",
    "competencies", "competencies and skills", "skills", "tools", "tooling", "frameworks",
    "libraries", "databases", "database", "platforms", "platform", "others", "other",
    "additional", "miscellaneous", "areas of expertise", "expertise", "knowledge", "areas",
}
_LEVEL_WORDS = {
    "advanced", "expert", "intermediate", "basic", "beginner", "fluent", "native",
    "professional", "proficient", "elementary", "conversational",
}
_LEVEL_SUFFIX_RE = re.compile(
    r"\s*[-–—:]\s*(?:\(?\b(?:advanced|expert|intermediate|basic|beginner|fluent|native|"
    r"professional|proficient|elementary|conversational)\b\)?)\s*$",
    re.IGNORECASE,
)
_SKILL_SPLIT_RE = re.compile(r"[,;|•·]+")


def _clean(line: str) -> str:
    text = _STRIP_MARKERS_RE.sub("", line).strip()
    return text.rstrip(":").strip()


def _is_bullet(line: str) -> bool:
    return bool(_BULLET_RE.match(line))


def _unique(items: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for item in items:
        low = item.lower()
        if low not in seen:
            seen.add(low)
            out.append(item)
    return out


def _range(match: re.Match[str]) -> tuple[str, str]:
    start = ((match.group(1) or "").strip() + " " + match.group(2)).strip()
    end = ((match.group(3) or "").strip() + " " + match.group(4)).strip()
    if match.group(4).lower() in ("present", "current", "now"):
        end = "present"
    return start, end


def _phone_number(text: str) -> str:
    for match in PHONE_RE.findall(text):
        digits = [ch for ch in match if ch.isdigit()]
        if len(digits) >= 7:
            return match.strip(" ()-.")
    return ""


def extract_cv_text(path: str | Path) -> str:
    """Read a CV file and return its plain text (paragraphs joined by newlines)."""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        try:
            from docx import Document

            document = Document(str(path))
        except Exception as exc:
            raise ValueError(f"cannot read the Word document: {exc}") from exc
        chunks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    chunks.append(" | ".join(cells))
        return "\n".join(chunks)
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"unsupported file type “{path.suffix or 'none'}” — "
                         f"use {', '.join(SUPPORTED_SUFFIXES)}")
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise ValueError(f"cannot read {path.name}: {exc.strerror or exc}") from exc
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return data.decode("latin-1", errors="replace")


def _is_section_heading(line: str) -> str | None:
    probe = _clean(line)
    if not probe or ":" in probe or "," in probe:
        return None
    return next((key for key, pattern in _SECTION_PATTERNS.items()
                 if pattern.search(probe.lower()) and len(probe) < 50), None)


def _inline_label_key(line: str) -> str | None:
    """Recognise a self-contained "Label: content" opener.

    ``Technologies: Python, SQL`` carries its own content, so the whole line
    joins the section without opening a flowing block (a bare ``Skills`` heading
    is what opens a multi-line section instead).
    """
    probe = _clean(line)
    if not probe or ":" not in probe:
        return None
    head = probe.split(":", 1)[0].strip().lower()
    if not head or len(head) > 40:
        return None
    for key, pattern in _SECTION_PATTERNS.items():
        if pattern.fullmatch(head):
            return key
    return None


def _split_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in lines:
        name = _is_section_heading(line)
        if name:
            current = name
            sections.setdefault(current, [])
            continue
        inline = _inline_label_key(line)
        if inline:
            sections.setdefault(inline, []).append(line)
        else:
            sections[current].append(line)
    return sections


def _detect_name(headers: list[str]) -> str:
    for line in headers:
        text = _clean(line)
        if not text or EMAIL_RE.search(text) or URL_RE.search(text):
            continue
        if text.strip("+() .-0123456789") == "" and any(ch.isdigit() for ch in text):
            continue
        if text.lower() in {"curriculum vitae", "resume", "résumé", "cv"}:
            continue
        if _NAME_RE.match(text) and ":" not in text:
            return text
    return ""


def _detect_headline(headers: list[str], name: str) -> str:
    after = False
    for line in headers:
        text = _clean(line)
        if not text:
            continue
        if after:
            if EMAIL_RE.search(text) or PHONE_RE.search(text) or URL_RE.search(text):
                continue
            if text.lower() in {"curriculum vitae", "resume", "résumé", "cv"}:
                continue
            if len(text) <= 120:
                return text
        elif text and text == name:
            after = True
    return ""


def _detect_location(headers: list[str]) -> str:
    for line in headers:
        text = _clean(line)
        match = _LOCATION_LABEL_RE.search(text)
        if match:
            return match.group(1).strip()
    for line in headers:
        text = _clean(line)
        if text and _CITY_RE.match(text):
            return text
    joined = " ".join(_clean(line) for line in headers)
    match = re.search(r"[A-ZÀ-Ž][A-Za-zÀ-ž .'\-]{2,},\s*[A-ZÀ-Ž][A-Za-zÀ-ž.'\-]{2,}", joined)
    return match.group(0).strip() if match else ""


def _detect_links(headers: list[str]) -> str:
    text = " ".join(headers)
    found = [m.group(0).rstrip(".,;)") for m in LINK_RE.finditer(text)]
    for match in URL_RE.finditer(text):
        url = match.group(0).rstrip(".,;)")
        if ("linkedin" in url.lower() or "github" in url.lower()) and url not in found:
            found.append(url)
    return " | ".join(_unique(found))


def _trim_summary(text: str, limit: int = 320) -> str:
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    head = text[:limit]
    cut = max(head.rfind(". "), head.rfind("! "), head.rfind("? "), head.rfind("; "))
    if cut > 160:
        return head[:cut + 1].strip()
    space = head.rfind(" ")
    return head[:space].strip() + " …" if space > 0 else head


def _strip_summary_label(line: str) -> str:
    """Drop a leading section label ("Professional Summary: …") from a summary line."""
    text = _clean(line)
    match = _SECTION_PATTERNS["summary"].match(text)
    if match and match.end() < len(text):
        rest = text[match.end():].lstrip(": \t").strip()
        if rest:
            return rest
    return text


def _detect_summary(lines: list[str], sections: dict[str, list[str]]) -> str:
    section = sections.get("summary")
    if section and any(_clean(line) for line in section):
        cleaned = [_strip_summary_label(line) for line in section if _clean(line)]
        return _trim_summary(" ".join(cleaned))
    hay = [line.strip() for line in lines if line.strip()]
    bound = next((i for i, line in enumerate(hay) if _is_section_heading(line)), len(hay))
    for i in range(bound):
        text = _clean(hay[i])
        if len(text) < 40:
            continue
        if EMAIL_RE.search(text) or PHONE_RE.search(text) or URL_RE.search(text):
            continue
        if _CITY_RE.match(text):
            continue
        if RANGE_RE.search(text) or _SINGLE_DATE_RE.match(text):
            continue
        if re.search(r"(?:—|–| \| | @ )", text):
            continue
        return _trim_summary(text)
    return ""


def _accept_skill(token: str) -> bool:
    low = token.lower()
    words = low.split()
    rejected = (len(token) < 2 or token.isdigit() or low in GENERIC_TERMS
                or low in STOPWORDS
                or bool(words) and all(w in GENERIC_TERMS or w in STOPWORDS for w in words))
    return not rejected


def _strip_skill_lead(token: str) -> str:
    match = re.match(
        r"^(?:proficient|skilled|experienced|fluent|knowledgeable|worked|comfortable|"
        r"familiar|focused|specialized|specialised|expert)\s+(?:with|in|using|on)\s+",
        token,
        re.IGNORECASE)
    return token[match.end():].strip() if match else token


def _parse_skills(lines: list[str]) -> list[str]:
    tokens: list[str] = []

    def collect(text: str) -> None:
        if not text or len(text) > 300:
            return
        parts = re.split(r"\s*:\s*", text)
        kept: list[str] = []
        for index, part in enumerate(parts):
            piece = part.strip()
            if not piece:
                continue
            if index == 0 and piece.lower() in _SKILL_LABELS:
                continue
            if index < len(parts) - 1:
                # A label word immediately before a colon ("SQL; Databases: "
                # "PostgreSQL") is a category heading: drop that trailing word.
                words = piece.split()
                if len(words) > 1 and words[-1].lower() in _SKILL_LABELS:
                    piece = " ".join(words[:-1]).strip()
            if piece:
                kept.append(piece)
        if not kept:
            return
        for piece in _SKILL_SPLIT_RE.split(" ".join(kept)):
            token = piece.strip()
            if not token:
                continue
            token = re.sub(r"\s*\([^)]*\)\s*$", "", token).strip()
            token = _LEVEL_SUFFIX_RE.sub("", token).strip()
            token = _strip_skill_lead(token)
            token = re.sub(r"^and\s+", "", token, flags=re.IGNORECASE).strip(" .,;:-")
            if token.lower() in _SKILL_LABELS:
                # "SQL; Databases: PostgreSQL" -> "PostgreSQL": a label word that
                # survives inside a compound line is a category heading, not a skill.
                continue
            if token and _accept_skill(token):
                tokens.append(token)

    for line in lines:
        text = _clean(line)
        if not text:
            continue
        if len(text) > 300:
            # Long "Skill A • Skill B • Skill C" runs (single-line competencies
            # lists) are >300 chars and would otherwise be dropped whole.
            for chunk in _SKILL_SPLIT_RE.split(text):
                collect(chunk.strip())
            continue
        collect(text)
    out: list[str] = []
    for token in tokens:
        if not any(t.lower() == token.lower() for t in out):
            out.append(pretty_term(token))
    return out[:40]


_LANG_LABEL_RE = re.compile(r"^\s*(?:foreign|spoken)?\s*languages?\s*[:;,–—-]\s*",
                            re.IGNORECASE)


def _is_level(value: str) -> bool:
    low = value.lower()
    if re.fullmatch(r"[abc][12]", low, re.IGNORECASE):
        return True
    return (low in _LEVEL_WORDS
            or low in ("mother tongue", "working knowledge", "native speaker",
                       "working proficiency", "professional working proficiency",
                       "full professional proficiency", "limited working proficiency",
                       "elementary proficiency", "native or bilingual proficiency"))


def _level_code(value: str) -> str:
    low = value.lower()
    if re.fullmatch(r"[abc][12]", low, re.IGNORECASE):
        return value.upper()
    return value.strip().capitalize()


def _language_item(text: str) -> str:
    cleaned = re.sub(r"[()]", " ", text).strip()
    bits = [b for b in re.split(r"[\s:\-–—/•·]+", cleaned) if b]
    if not bits:
        return ""
    name, *rest = bits
    if name.lower() in _SKILL_LABELS or name.lower() in ("language", "languages"):
        name = rest[0] if rest else ""
        rest = rest[1:]
    if rest:
        level = " ".join(rest).strip()
        if _is_level(level):
            return f"{pretty_term(name)} ({_level_code(level)})"
        if len(rest) == 1:
            return pretty_term(name + " " + rest[0])
    if _is_level(name) and rest:
        return f"{pretty_term(' '.join(rest))} ({_level_code(name)})"
    return pretty_term(name)


def _parse_languages(lines: list[str]) -> str:
    items: list[str] = []
    for line in lines:
        text = _clean(line)
        if not text:
            continue
        if _LANG_LABEL_RE.match(text):
            text = _LANG_LABEL_RE.sub("", text)
        if not text:
            continue
        # Separate pipe-delimited chunks first so location/mobility fragments
        # ("Based in Nairobi, Kenya | Open to UAE/GCC …") are dropped whole
        # instead of leaking country names into the language list.
        chunks = [c.strip() for c in re.split(r"\s*\|\s*", text)]
        for chunk in chunks:
            if not chunk:
                continue
            low = chunk.lower()
            if low.startswith(("based in", "located", "open to", "willing", "mobile",
                               "relocation", "prefer")):
                continue
            for piece in re.split(r"\s*[;,]\s*", chunk):
                if not piece or _looks_like_location(piece):
                    continue
                item = _language_item(piece)
                if item and not any(existing.lower() == item.lower() for existing in items):
                    items.append(item)
    return ", ".join(items)


_CITY_NAMES = {
    "Amsterdam", "Athens", "Auckland", "Barcelona", "Beijing", "Bengaluru", "Berlin",
    "Brussels", "Budapest", "Buenos Aires", "Cairo", "Chicago", "Copenhagen", "Delhi",
    "Dubai", "Dublin", "Frankfurt", "Geneva", "Hamburg", "Helsinki", "Hong Kong",
    "Istanbul", "Jakarta", "Kuala Lumpur", "Lisbon", "London", "Los Angeles", "Madrid",
    "Manchester", "Melbourne", "Mexico City", "Miami", "Milan", "Montreal", "Moscow",
    "Munich", "Mumbai", "New York", "Oslo", "Paris", "Prague", "Rome", "San Francisco",
    "São Paulo", "Seoul", "Shanghai", "Singapore", "Stockholm", "Sydney", "Tel Aviv",
    "Tokyo", "Toronto", "Vienna", "Warsaw", "Washington", "Zurich", "Abidjan",
    "Accra", "Addis Ababa", "Nairobi", "Lagos", "Johannesburg", "Cape Town",
}


def _looks_like_location(text: str) -> bool:
    value = text.strip()
    if not value:
        return False
    low = value.lower()
    if "," in value or _CITY_RE.match(value):
        return True
    if value.title() in _CITY_NAMES:
        return True
    return low in {"remote", "hybrid", "onsite", "remote-first", "remote friendly", "on-site",
                   "in-office", "in office"}


_ROLE_START_RE = re.compile(
    r"^(?:it(?:\s+support)?|software|hardware|computer|network|sales|business|data|"
    r"finance|financial|operations?|logistics|supply|marketing|support|engineering|"
    r"engineer|developer|qa|quality|admin|administrative|customer|project|product|"
    r"procurement|inventory|transport|warehouse|analyst|manager|coordinator|supervisor|"
    r"specialist|technician)\b",
    re.IGNORECASE,
)


def _split_role_line(text: str) -> tuple[str, str, str]:
    label = _clean(text)
    pieces = [p.strip() for p in label.split("|") if p.strip()]
    if not pieces:
        return "", "", ""
    title = company = location = ""
    head = pieces[0]
    for sep in (" — ", " – ", " - ", " at ", "@"):
        if sep in head:
            title, company = (part.strip() for part in head.split(sep, 1))
            # "DHL - IT Support" is Company - Role while the common pattern is
            # "Role - Company": flip when the piece after the dash is role-like.
            if sep == " - " and _ROLE_START_RE.match(company):
                title, company = company, title
            break
    else:
        title = head
    rest = pieces[1:]
    if rest:
        tail = rest[-1]
        has_company = bool(company)
        three_part = len(rest) == 2 and not _looks_like_location(rest[0])
        if _looks_like_location(tail) or (has_company and len(rest) == 1) or three_part:
            location = tail
            leftover = rest[:-1]
        else:
            leftover = rest
        if leftover:
            if not company and leftover[0]:
                company = leftover[0]
                if len(leftover) > 1:
                    company += " " + " ".join(leftover[1:])
            elif company and leftover:
                company += " " + " ".join(leftover)
    return title.strip(), company.strip(), location.strip()


def _date_of(line: str) -> tuple[str, str] | None:
    """Return ``(start, end)`` when the line carries a date range or a single year."""
    match = RANGE_RE.search(line)
    if match:
        return _range(match)
    match = _SINGLE_DATE_RE.match(line)
    if match:
        return match.group(0).strip(), ""
    return None


_ROLE_SEP_RE = re.compile(r"\s(?:–|—)\s|\s at \s|\s-\s")


def _is_next_header(line: str) -> bool:
    """True when a line inside a position block opens the *next* role instead of
    being prose, so dense CVs keep every job separate (headers no longer leak
    into the previous entry's summary)."""
    text = line.strip()
    if not text or len(text) > 200 or _is_bullet(text) or _date_of(text):
        return False
    if "|" in text:
        return True
    if text.rstrip().endswith((".", "!", "?", ";", ":")) or "," in text:
        return False
    return bool(_ROLE_SEP_RE.search(text))


def _date_tail(line: str, match: re.Match[str]) -> tuple[str, list[str]]:
    """Split the text after a date range on a date line into ``(location, bullets)``."""
    rest = line[match.end():].strip().lstrip("|").strip()
    if not rest:
        return "", []
    parts = [part.strip() for part in rest.split("•")]
    location = parts[0] if parts and _looks_like_location(parts[0]) else ""
    return location, [part for part in parts[1:] if part]


def _parse_experience(lines: list[str]) -> list[ExperienceEntry]:
    entries: list[ExperienceEntry] = []
    pending_header: list[str] = []
    pending_bullets: list[str] = []
    active: dict[str, object] | None = None

    def flush() -> None:
        nonlocal active, pending_header, pending_bullets
        if active is None:
            return
        title, company, location = _split_role_line(" | ".join(active["header"]))
        start, end = active["dates"]
        body = active["body"]
        bullets = [_clean(line) for line in body if _is_bullet(line)]
        prose = [_clean(line) for line in body if not _is_bullet(line) and _clean(line)]
        location = active["location"] or location
        if title or company or bullets or prose:
            entries.append(ExperienceEntry(
                title=title, company=company, location=location,
                start=start, end=end, summary=" ".join(prose),
                highlights=_unique(pending_bullets + bullets)))
        active = None
        pending_header.clear()
        pending_bullets.clear()

    for raw in lines:
        line = raw.strip()
        if not line:
            flush()
            continue
        if re.match(r"^(?:earlier|additional|other|relevant|volunteer|academic|"
                    r"teaching|freelance|consulting|professional)\s+experience\b",
                    line, re.IGNORECASE):
            pending_header.clear()
            continue
        dates = _date_of(line)
        if dates is not None:
            # A new position starts at every date line, so dense CVs (headers
            # carrying "Role | Company | 2021 - Present" with no blank lines)
            # split into one entry per job instead of swallowing the extras.
            match = RANGE_RE.search(line)
            prefix = line[:match.start()].strip() if match else ""
            location, trailing_bullets = _date_tail(line, match) if match else ("", [])
            # Capture the queued header *before* flush(), which clears it.
            header = list(pending_header)
            flush()
            if prefix:
                header.append(prefix)
            if trailing_bullets:
                pending_bullets.extend(trailing_bullets)
            active = {"header": [h for h in header if h],
                      "dates": dates, "body": [], "location": location}
            pending_header.clear()
            continue
        if active is not None:
            if _is_next_header(line):
                pending_header.append(_clean(line))
            else:
                active["body"].append(line)
        elif _is_bullet(raw):
            pending_bullets.append(_clean(line))
        else:
            pending_header.append(_clean(line))
    flush()
    if pending_header:
        title, company, location = _split_role_line(" | ".join(pending_header))
        text = (" ".join(pending_header)).lower()
        denied = any(term in text for term in ("references", "certification",
                                               "available on request", "interests",
                                               "hobbies", "additional information"))
        words = " ".join(pending_header).split()
        if (title or company or pending_bullets) and not denied and len(words) <= 12:
            entries.append(ExperienceEntry(title=title, company=company, location=location,
                                           start="", end="",
                                           highlights=list(pending_bullets)))
    return entries


def _assign_edu(text: str, entry: EducationEntry) -> None:
    if not text:
        return
    for sep in (" — ", " – ", " | ", ","):
        if sep in text:
            a, b = (part.strip() for part in text.split(sep, 1))
            if _DEGREE_RE.search(a):
                a = re.sub(r"\s*[-–—]\s*(?:in progress|ongoing|current)\s*$", "", a,
                           flags=re.IGNORECASE).strip()
                entry.degree = entry.degree or a
                entry.school = entry.school or b
                return
    if _DEGREE_RE.search(text):
        entry.degree = entry.degree or text
        return
    if _SCHOOL_RE.search(text):
        entry.school = entry.school or text
        return
    if not entry.degree:
        entry.degree = text
    elif not entry.school:
        entry.school = text
    else:
        entry.details = (entry.details + "\n" + text).strip()


def _parse_education(lines: list[str]) -> list[EducationEntry]:
    entries: list[EducationEntry] = []
    current: EducationEntry | None = None

    def flush() -> None:
        nonlocal current
        if current is not None and (current.degree or current.school):
            entries.append(current)
        current = None

    for raw in lines:
        line = raw.strip()
        if not line:
            flush()
            continue
        if _DEGREE_RE.search(line) and current is not None and current.degree:
            # A fresh "Degree — School [dates]" line starts its own entry even in
            # dense CVs where the lines are not separated by blank lines.
            flush()
        if current is None:
            current = EducationEntry()
        match = RANGE_RE.search(line)
        if match:
            current.start, current.end = _range(match)
            head = line[:match.start()].rstrip(" |\t").strip()
            if head:
                _assign_edu(head, current)
            tail = line[match.end():].lstrip(" |\t").strip()
            if tail:
                _assign_edu(tail, current)
            continue
        single = _SINGLE_DATE_RE.match(line)
        if single:
            current.start = single.group(0).strip()
            continue
        _assign_edu(_clean(line), current)
    flush()
    return entries


def parse_cv(path: str | Path) -> Profile:
    """Parse a CV file into a Profile draft (empty fields simply stay empty)."""
    path = Path(path)
    text = extract_cv_text(path)
    if not text.strip():
        raise ValueError(f"no usable text found in {path.name}")
    sections = _split_sections(text.splitlines())
    headers = sections.get("header") or []
    joined = " ".join(headers)
    name = _detect_name(headers)
    headline = _detect_headline(headers, name)
    # The headline line (e.g. "SUPPLY CHAIN | OPERATIONS | DATA REPORTING") and the
    # name are not locations: drop both so their words cannot leak into the location.
    effective = [line for line in headers
                 if _clean(line) not in (headline, name)]
    seniority = next((word for word in ("junior", "mid", "senior", "lead", "principal")
                      if word in headline.lower()), "")
    mobility = f"{joined} {' '.join(sections.get('languages') or [])}".lower()
    profile = Profile(
        full_name=name,
        headline=headline,
        email=next(iter(EMAIL_RE.findall(joined)), ""),
        phone=_phone_number(joined),
        location=_detect_location(effective),
        links=_detect_links(headers),
        summary=_detect_summary(text.splitlines(), sections),
        skills=_parse_skills(sections.get("skills") or []),
        languages=_parse_languages(sections.get("languages") or []),
        experience=_parse_experience(sections.get("experience") or []),
        education=_parse_education(sections.get("education") or []),
        willing_to_relocate=bool(
            re.search(r"\bopen\s+to\s+(?:relocat\w+|uae\b|gcc\b|international|the uae\b)",
                      mobility)),
        seniority=seniority,
    )
    return profile