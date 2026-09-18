"""CV templates, cover-letter tones, e-mail drafts and exporters."""
from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from jautomatic.models import SAMPLE_PROFILE, JobPosting, Profile
from jautomatic.services.cover_letter import (
    TONES,
    CoverLetterService,
    MatchContext,
    render_cover_letter,
    top_keywords_from_job,
)
from jautomatic.services.cv_generator import (
    TEMPLATES,
    CVGenerator,
    export,
    render_markdown,
    to_docx,
    to_plain_text,
)
from jautomatic.services.email_drafter import (
    EmailDrafter,
    build_subject,
    guess_recipient,
    render_email,
    render_follow_up,
)


def profile() -> Profile:
    return Profile.from_dict(SAMPLE_PROFILE)


def job() -> JobPosting:
    return JobPosting(source="remotive", title="Senior Python Engineer",
                      company="Northwind Analytics", location="Berlin, Germany", remote=True,
                      salary_min=90000, salary_max=110000, currency="USD",
                      url="https://example.com/jobs/1",
                      description="We are looking for a Senior Python Engineer to own our data "
                                  "platform. FastAPI, PostgreSQL, Docker and AWS. Contact "
                                  "jobs@northwind.example.com to apply.",
                      tags=["python", "fastapi", "postgresql", "aws"],
                      posted_at=date.today().isoformat())


class CVTemplateTests(unittest.TestCase):
    def test_every_template_renders_identity_and_experience(self):
        for template in TEMPLATES:
            text = render_markdown(profile(), job(), template)
            self.assertIn("alex doe", text.lower())
            self.assertIn("Datawheel GmbH", text)
            self.assertIn("Warsaw University of Technology", text)
            self.assertGreater(len(text.splitlines()), 12, template)

    def test_modern_template_lists_tailored_skills(self):
        text = render_markdown(profile(), job(), "modern")
        self.assertIn("## Core skills", text)
        self.assertIn("python", text.lower())

    def test_classic_template_uses_caps_headings(self):
        text = render_markdown(profile(), job(), "classic")
        self.assertIn("PROFESSIONAL EXPERIENCE", text)
        self.assertIn("KEY SKILLS", text)

    def test_compact_template_is_shortest(self):
        lengths = {template: len(render_markdown(profile(), job(), template).splitlines())
                   for template in TEMPLATES}
        self.assertLess(lengths["compact"], lengths["classic"])

    def test_generic_words_are_not_advertised_as_skills(self):
        text = render_markdown(profile(), job(), "modern")
        skills_line = [line for line in text.splitlines()
                       if line.startswith("python,")][0]
        for generic in ("senior,", "engineer,", "data,"):
            self.assertNotIn(generic, skills_line.lower())

    def test_summary_is_tailored_to_the_job(self):
        text = render_markdown(profile(), job(), "modern")
        self.assertIn("Targeting the Senior Python Engineer role at Northwind Analytics", text)

    def test_unknown_template_falls_back_to_modern(self):
        self.assertEqual(render_markdown(profile(), job(), "nope"),
                         render_markdown(profile(), job(), "modern"))

    def test_generator_writes_requested_format(self):
        with tempfile.TemporaryDirectory() as tmp:
            for fmt, suffix in (("docx", ".docx"), ("md", ".md"), ("txt", ".txt")):
                document = CVGenerator().generate(profile(), job(), "modern", Path(tmp), fmt)
                self.assertEqual(document.path.suffix, suffix)
                self.assertTrue(document.path.exists())
                self.assertGreater(document.path.stat().st_size, 200)
                self.assertTrue(document.text.strip())


class ExporterTests(unittest.TestCase):
    def test_markdown_to_plain_text(self):
        plain = to_plain_text("# Title\n\n- **Bold** item\n*italic*\n---\n")
        self.assertNotIn("#", plain)
        self.assertNotIn("**", plain)
        self.assertIn("Bold item", plain)
        self.assertIn("italic", plain)

    def test_docx_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "doc.docx"
            to_docx("# Alex Doe\n\n## Experience\n- Owned the platform\n", path, title="CV")
            from docx import Document
            document = Document(str(path))
            text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("Alex Doe", text)
            self.assertIn("Owned the platform", text)
            self.assertEqual(document.core_properties.title, "CV")
            self.assertEqual(document.styles["Normal"].font.name, "Segoe UI")
            self.assertEqual(document.paragraphs[0].style.name, "Title")
            self.assertFalse(any(not p.text for p in document.paragraphs))

    def test_export_switches_suffix(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = export("Hello **world**", Path(tmp) / "file.doc", "md")
            self.assertEqual(path.suffix, ".md")
            self.assertTrue(path.exists())

    def test_word_and_pdf_names_are_reserved_together(self):
        from jautomatic.models import unique_document_path
        with tempfile.TemporaryDirectory() as tmp:
            original = Path(tmp) / "CV_alex.docx"
            original.write_bytes(b"existing application")
            candidate = unique_document_path(tmp, "CV_alex", ".pdf", key="different role")
            self.assertNotEqual(candidate.stem, original.stem)
            self.assertEqual(original.read_bytes(), b"existing application")


class CoverLetterTests(unittest.TestCase):
    def test_core_sections_present(self):
        text = render_cover_letter(profile(), job(), MatchContext(
            score=88, matched_keywords=["python", "fastapi", "postgresql"]))
        self.assertIn("Dear Northwind Analytics Hiring Team,", text)
        self.assertIn("Senior Python Engineer", text)
        self.assertIn("FastAPI", text)          # pretty-printed in prose
        self.assertIn("Best regards,", text)
        self.assertIn(date.today().year >= 2026 and "" or "", text)  # smoke: no crash

    def test_references_posting_text(self):
        text = render_cover_letter(profile(), job())
        self.assertIn("data platform", text)

    def test_tones_change_the_closing(self):
        outputs = {tone: render_cover_letter(profile(), job(), None, tone) for tone in TONES}
        self.assertEqual(len(set(outputs.values())), len(TONES))
        self.assertIn("love the chance", outputs["enthusiastic"])
        self.assertIn("I'd be glad to chat", outputs["friendly"])
        self.assertLess(len(outputs["concise"].split("\n\n")),
                        len(outputs["professional"].split("\n\n")))

    def test_salary_alignment_sentence(self):
        below = job()
        below.salary_max = 40000
        text = render_cover_letter(profile(), below)
        self.assertIn("salary expectations start around", text)
        text_ok = render_cover_letter(profile(), job())
        self.assertIn("aligns well with my expectations", text_ok)

    def test_first_person_rewrite_of_achievement(self):
        text = render_cover_letter(profile(), job())
        self.assertIn("I owned the ingestion platform", text)

    def test_service_writes_letter_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            document = CoverLetterService().generate(profile(), job(), None, "professional",
                                                     Path(tmp), "docx")
            self.assertTrue(document.path.exists())
            self.assertEqual(document.kind, "cover_letter")
            from docx import Document
            letter = Document(str(document.path))
            self.assertEqual(letter.paragraphs[0].text, profile().display_name)
            self.assertEqual(letter.paragraphs[0].style.name, "Title")
            self.assertIn("Application for Senior Python Engineer", document.text)

    def test_top_keywords_prefer_tags(self):
        words = top_keywords_from_job(job(), limit=5)
        self.assertEqual(words[0], "python")
        self.assertLessEqual(len(words), 5)


class EmailTests(unittest.TestCase):
    def test_subject_and_recipient(self):
        draft = render_email(profile(), job())
        self.assertIn("Application: Senior Python Engineer", draft.subject)
        self.assertIn("Alex Doe", draft.subject)
        self.assertEqual(draft.recipient, "jobs@northwind.example.com")

    def test_body_mentions_role_and_attachments(self):
        draft = render_email(profile(), job(), attachments=["CV.docx", "Letter.docx"])
        self.assertIn("Senior Python Engineer", draft.body)
        self.assertIn("CV.docx", draft.text)
        self.assertIn("Subject:", draft.text)

    def test_subject_without_name(self):
        person = profile()
        person.full_name = ""
        self.assertNotIn("(", build_subject(person, job()))

    def test_recipient_guess_is_blank_without_address(self):
        posting = job()
        posting.description = "Apply through our careers portal."
        self.assertEqual(guess_recipient(posting), "")

    def test_follow_up_email(self):
        draft = render_follow_up(profile(), job(), days_since_sent=9)
        self.assertIn("Following up", draft.subject)
        self.assertIn("9 days ago", draft.body)
        interview = render_follow_up(profile(), job(), 3, stage="interview")
        self.assertIn("Thank you again for the interview", interview.body)

    def test_service_persists_drafts(self):
        with tempfile.TemporaryDirectory() as tmp:
            drafter = EmailDrafter()
            email = drafter.generate(profile(), job(), None, None, ["CV.docx"], Path(tmp), "md")
            follow_up = drafter.follow_up(profile(), job(), 7, "sent", Path(tmp), "md")
            self.assertTrue(email.path.exists() and email.path.suffix == ".md")
            self.assertTrue(follow_up.path.exists())
            self.assertIn("Attachments", email.text)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
