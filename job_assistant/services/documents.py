"""Truthful, ATS-friendly application document generation."""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from job_assistant.models.entities import CandidateProfile


def safe_name(value: str, fallback: str = "Document") -> str:
    value = re.sub(r"[^A-Za-z0-9._ -]+", "", value).strip(" .")
    return re.sub(r"\s+", "_", value)[:80] or fallback


class DocumentGenerator:
    def __init__(self, applications_dir: str | Path):
        self.applications_dir = Path(applications_dir)

    def prepare(self, profile: CandidateProfile, job: dict[str, Any], analysis: dict[str, Any]) -> dict[str, str]:
        company = safe_name(str(job.get("company") or "Company"))
        title = safe_name(str(job.get("title") or "Position"))
        year = datetime.now().strftime("%Y")
        folder = self.applications_dir / year / company / title
        folder.mkdir(parents=True, exist_ok=True)
        stem = f"{company}_{title}"
        job_details = folder / "Job_Details.txt"
        job_details.write_text(self._job_details(job), encoding="utf-8")
        (folder / "Original_Job_Description.txt").write_text(str(job.get("description") or ""), encoding="utf-8")
        match_text = self._match_analysis(analysis)
        (folder / "Match_Analysis.txt").write_text(match_text, encoding="utf-8")
        self._write_pdf(folder / "Original_Job_Description.pdf", "Original Job Description", [str(job.get("description") or "No description supplied")])
        self._write_pdf(folder / "Match_Analysis.pdf", "Match Analysis", match_text.splitlines())

        cv_body = self._cv_sections(profile, job, analysis)
        cv_docx = folder / f"CV_{stem}.docx"
        cv_pdf = folder / f"CV_{stem}.pdf"
        self._write_docx(cv_docx, profile.full_name or "Candidate", cv_body)
        self._write_pdf(cv_pdf, profile.full_name or "Candidate CV", cv_body)

        letter_body = self._cover_letter(profile, job, analysis)
        letter_docx = folder / f"Cover_Letter_{stem}.docx"
        letter_pdf = folder / f"Cover_Letter_{stem}.pdf"
        self._write_docx(letter_docx, "Cover Letter", [("", letter_body)])
        self._write_pdf(letter_pdf, "Cover Letter", letter_body.splitlines())

        record = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "job": {key: job.get(key, "") for key in ("title", "company", "location", "job_url", "source_name")},
            "match_analysis": analysis,
            "documents": {
                "cv_docx": str(cv_docx), "cv_pdf": str(cv_pdf),
                "cover_letter_docx": str(letter_docx), "cover_letter_pdf": str(letter_pdf),
            },
        }
        (folder / "Application_Record.json").write_text(json.dumps(record, indent=2), encoding="utf-8")
        return {"folder": str(folder), "cv_docx": str(cv_docx), "cv_pdf": str(cv_pdf), "cover_letter_docx": str(letter_docx), "cover_letter_pdf": str(letter_pdf)}

    @staticmethod
    def _job_details(job: dict[str, Any]) -> str:
        return "\n".join([
            f"Position: {job.get('title', '')}", f"Company: {job.get('company', '')}",
            f"Location: {job.get('location', '')}", f"Source: {job.get('source_name', '')}",
            f"Job URL: {job.get('job_url', '')}", f"Posted: {job.get('posted_at', '')}",
            f"Closing: {job.get('closing_at', '')}", "", str(job.get("description") or "No description supplied"),
        ])

    @staticmethod
    def _match_analysis(analysis: dict[str, Any]) -> str:
        return "\n".join([
            f"Match score: {analysis.get('score', 0)}%", f"Category: {analysis.get('category', '')}",
            f"Recommendation: {analysis.get('recommendation', '')}", "", "Reasons:",
            *[f"- {item}" for item in analysis.get("reasons", [])], "", "Matched skills:",
            *[f"- {item}" for item in analysis.get("matched", [])], "", "Missing mandatory evidence:",
            *[f"- {item}" for item in analysis.get("missing", [])], "", "Preferred skills not evidenced:",
            *[f"- {item}" for item in analysis.get("preferred", [])],
        ])

    @staticmethod
    def _cv_sections(profile: CandidateProfile, job: dict[str, Any], analysis: dict[str, Any]) -> list[tuple[str, str]]:
        matched = analysis.get("matched", [])
        skills = matched + [item for item in profile.technical_skills + profile.software_skills if item not in matched]
        sections: list[tuple[str, str]] = [("Professional Summary", profile.summary or profile.headline)]
        sections.append(("Core Skills", " • ".join(skills)))
        if profile.experience:
            sections.append(("Professional Experience", profile.experience))
        if profile.achievements:
            sections.append(("Selected Achievements", profile.achievements))
        if profile.education:
            sections.append(("Education", profile.education))
        if profile.qualifications:
            sections.append(("Qualifications", profile.qualifications))
        if profile.certifications:
            sections.append(("Certifications", profile.certifications))
        return [(heading, value) for heading, value in sections if value.strip()]

    @staticmethod
    def _cover_letter(profile: CandidateProfile, job: dict[str, Any], analysis: dict[str, Any]) -> str:
        title = str(job.get("title") or "the advertised position")
        company = str(job.get("company") or "your organisation")
        name = profile.full_name or "the candidate"
        matched = ", ".join(analysis.get("matched", [])[:4]) or "the relevant experience outlined in my CV"
        paragraphs = [
            f"Dear Hiring Manager,",
            f"I am writing to apply for the {title} position at {company}. My background includes {matched}, and I am interested in bringing that experience to this opportunity.",
            f"The role is a strong match for the evidence in my profile, particularly in relation to {matched}. I would welcome the opportunity to discuss how my experience, qualifications, and practical approach could contribute to your team.",
        ]
        if profile.achievements:
            paragraphs.append(f"My relevant achievements include {profile.achievements[:600].rstrip('.')}.")
        paragraphs.extend(["Thank you for reviewing my application. I would be pleased to provide any additional information required.", f"Kind regards,\n{name}"])
        return "\n\n".join(paragraphs)

    @staticmethod
    def _write_docx(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("python-docx is required to generate DOCX files") from exc
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        document.add_heading(title, level=0)
        for heading, body in sections:
            if heading:
                document.add_heading(heading, level=1)
            for paragraph in body.split("\n"):
                document.add_paragraph(paragraph.strip())
        for style in document.styles:
            if style.name == "Normal":
                style.font.name = "Arial"
                style.font.size = Pt(10)
        document.save(path)

    @staticmethod
    def _write_pdf(path: Path, title: str, lines: list[str]) -> None:
        try:
            from reportlab.lib.enums import TA_LEFT
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import mm
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError as exc:
            raise RuntimeError("reportlab is required to generate PDF files") from exc
        styles = getSampleStyleSheet()
        body = ParagraphStyle("AssistantBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, alignment=TA_LEFT, spaceAfter=5)
        heading = ParagraphStyle("AssistantHeading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, spaceAfter=10)
        story = [Paragraph(title, heading)]
        for line in lines:
            safe = str(line).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe or "&nbsp;", body))
        SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=15 * mm, bottomMargin=15 * mm).build(story)
